#!/usr/bin/env python3
"""Reproducibly inspect/import the newly uploaded source without claiming it is raw counts."""
import csv, gzip, hashlib, json, pathlib, sys, zipfile
sys.path.insert(0, '/workspace/cloud-setup/python-module1')
import libarchive
import openpyxl

root=pathlib.Path('/workspace/GSE194331-RNAseq-R')
source=root/'data/module1_inputs/GSE194331.rar'
extracted=root/'data/module1_inputs/extracted/GSE194331.csv'
derived=root/'data/module1_inputs/derived/GSE194331_original_values.tsv.gz'
derived.parent.mkdir(parents=True,exist_ok=True)
extracted.parent.mkdir(parents=True,exist_ok=True)
source_sha=hashlib.sha256(source.read_bytes()).hexdigest()
provenance_path=root/'results/module1_audit/original_source_provenance.json'
prior=json.loads(provenance_path.read_text()) if provenance_path.exists() else None
if prior and prior.get('source_sha256')!=source_sha:
    raise RuntimeError('Uploaded archive changed since the recorded extraction; inspect before proceeding.')
members=[]
with libarchive.file_reader(str(source)) as archive:
    for entry in archive:
        members.append({'name':entry.pathname,'size_bytes':entry.size})
        if entry.pathname != 'GSE194331.csv':
            raise RuntimeError('Unexpected archive member; inspect before extraction.')
        if not extracted.exists():
            with extracted.open('xb') as stream:
                for block in entry.get_blocks(): stream.write(block)
if not zipfile.is_zipfile(extracted):
    raise RuntimeError('Expected inspected OOXML workbook; re-inspect file type.')
extracted_sha=hashlib.sha256(extracted.read_bytes()).hexdigest()
if prior and prior.get('extracted_sha256') and prior.get('extracted_sha256')!=extracted_sha:
    raise RuntimeError('Previously extracted workbook differs from recorded source; do not overwrite.')
with zipfile.ZipFile(extracted) as z:
    if 'xl/workbook.xml' not in z.namelist():raise RuntimeError('Not an OOXML workbook.')
tmp=derived.with_name(derived.name+'.partial')
types=set();rows=0
with extracted.open('rb') as stream, gzip.open(tmp,'wt',encoding='utf-8',newline='') as out:
    workbook=openpyxl.load_workbook(stream,read_only=True,data_only=True)
    if workbook.sheetnames != ['Sheet1']:raise RuntimeError('Unexpected workbook sheets.')
    sheet=workbook['Sheet1']
    writer=csv.writer(out,delimiter='\t')
    for i,row in enumerate(sheet.iter_rows(values_only=True)):
        if i == 0:
            header=list(row)
            if len(header)!=120 or header[0]!='id':raise RuntimeError('Unexpected expression table header.')
            writer.writerow(row)
        else:
            types.add(type(row[0]).__name__)
            writer.writerow(row);rows+=1
            if rows%10000==0:print('Imported source rows:',rows,flush=True)
    workbook.close()
if derived.exists():
    # Gzip timestamps may differ: compare decompressed content, not container hashes.
    def content_hash(p):
        h=hashlib.sha256()
        with gzip.open(p,'rb') as f:
            for block in iter(lambda:f.read(1024*1024),b''):h.update(block)
        return h.hexdigest()
    if content_hash(derived)!=content_hash(tmp):raise RuntimeError('Existing derived values differ; do not overwrite.')
    tmp.unlink()
else:tmp.rename(derived)
provenance=dict(source_archive=str(source),source_sha256=source_sha,
    archive_members=members,archive_source='origin GitHub commit 94abd7389e06991243a8c467c0dcedab5e0ea40f, GSE194331.rar',
    extraction_file=str(extracted),extracted_sha256=extracted_sha,actual_format='OOXML Excel workbook despite .csv extension',
    derived_file=str(derived),n_genes=rows,n_samples=len(header)-1,gene_id_python_types=sorted(types),
    import_method='openpyxl read_only=True, data_only=True; unchanged numeric cell values; no rounding/filtering/normalization',
    expression_scale='UNKNOWN_CONTINUOUS_SOURCE; original upload does not establish raw-count provenance',
    library_versions=dict(openpyxl=openpyxl.__version__,libarchive_c='5.3'))
(root/'results/module1_audit/original_source_provenance.json').write_text(json.dumps(provenance,ensure_ascii=False,indent=2)+'\n')
print(f'Imported {rows} genes, {len(header)-1} samples. Source units remain unknown.',flush=True)
